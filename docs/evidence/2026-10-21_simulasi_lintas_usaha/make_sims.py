"""Four business types as Mankeu template v3 files, plus independent expected figures.

The expected figures are computed here from first principles (profit = 0 at break-even), not by
calling the app, so a mismatch on screen is a real finding.

Historical record (2026-10-29, phase-1 cleanup P8): this run predates template v5 (DEC-028); the
current importer rejects these files by their headers. Kept unchanged as the evidence of that run;
the v5 generator is `docs/evidence/2026-10-22_simulasi_skala_umkm/make_sims.py`.
"""
from decimal import Decimal as D, ROUND_HALF_UP
import json, os
from openpyxl import Workbook

OUT = os.path.dirname(os.path.abspath(__file__))
TX = ["Tanggal", "Keterangan", "Kategori", "Pemasukan (Rp)", "Pengeluaran (Rp)", "Produk", "Jumlah", "Sifat Biaya", "Umur Pakai (bulan)"]
BUDGET = ["Bulan", "Kategori", "Item", "Jumlah", "Satuan", "Harga Satuan (Rp)", "Total (Rp)", "Berulang", "Sifat Biaya", "Umur Pakai (bulan)"]
PRODUCT = ["Nama", "Satuan", "Harga Jual (Rp)", "Target per Bulan"]
RECIPE = ["Produk", "Komponen", "Biaya (Rp)", "Untuk Berapa Unit"]
STOCK = ["Nama", "Satuan", "Harga Beli (Rp)", "Jumlah Stok", "Batas Minimum"]

JUAL, LAIN_IN = "Penjualan & Jasa", "Pendapatan Lain"
BAHAN, GAJI, SEWA, LISTRIK, TRANS, LAIN = "Bahan Baku & Stok", "Gaji & Upah", "Sewa Tempat", "Listrik, Air & Internet", "Transportasi", "Biaya Lain-lain"
MODAL, ASET = "Setoran Modal Pemilik", "Beli Peralatan & Aset"
NATURE = {BAHAN: "V", TRANS: "V", GAJI: "F", SEWA: "F", LISTRIK: "F", LAIN: "F"}

def r0(x): return int(D(x).quantize(D(1), ROUND_HALF_UP))

SIMS = {
    "A_toko_kelontong": dict(
        jenis="Dagang (toko kelontong, 3 produk)",
        products=[("Beras 5 kg", "karung", 75000, 120, [("Harga beli", 65000, 1)]),
                  ("Minyak 2 L", "botol", 38000, 200, [("Harga beli", 33000, 1)]),
                  ("Gula 1 kg", "kg", 18000, 300, [("Harga beli", 15500, 1)])],
        plan=[(None, SEWA, "Sewa kios", 1500000, "Bulanan", None, None),
              (None, LISTRIK, "Listrik", 300000, "Bulanan", None, None),
              (None, GAJI, "Gaji penjaga", 2000000, "Bulanan", None, None),
              ("2026-11", ASET, "Rak dan freezer", 6000000, "Sekali", None, 48),
              ("2026-11", MODAL, "Modal awal", 30000000, "Sekali", None, None)],
        tx=[("2026-10-31", "Jual beras Oktober", JUAL, 110 * 75000, None, "Beras 5 kg", "110", None, None),
            ("2026-10-31", "Jual minyak Oktober", JUAL, 210 * 38000, None, "Minyak 2 L", "210", None, None),
            ("2026-10-31", "Jual gula Oktober", JUAL, 280 * 18000, None, "Gula 1 kg", "280", None, None),
            ("2026-10-02", "Kulakan barang", BAHAN, None, 18600000, None, None, None, None),
            ("2026-10-01", "Sewa kios", SEWA, None, 1500000, None, None, None, None),
            ("2026-10-25", "Listrik", LISTRIK, None, 320000, None, None, None, None),
            ("2026-10-30", "Gaji penjaga", GAJI, None, 2000000, None, None, None, None),
            ("2026-10-01", "Rak dan freezer", ASET, None, 6000000, None, None, None, 48),
            ("2026-10-01", "Modal awal", MODAL, 30000000, None, None, None, None, None)],
        stock=[("Beras 5 kg", "karung", 65000, "40", "10")],
    ),
    "B_laundry": dict(
        jenis="Jasa (laundry kiloan, 2 layanan, periode 12 bulan)",
        products=[("Cuci setrika", "kg", 8000, 1500, [("Deterjen dan pewangi", 45000, 30), ("Plastik", 500, 1), ("Listrik dan air", 700, 1)]),
                  ("Cuci kilat", "kg", 12000, 300, [("Deterjen dan pewangi", 45000, 30), ("Plastik", 500, 1), ("Listrik dan air", 700, 1)])],
        plan=[(None, GAJI, "Gaji 2 karyawan", 3600000, "Bulanan", None, None),
              (None, SEWA, "Sewa ruko", 2000000, "Bulanan", None, None),
              ("2026-11", LAIN, "Izin usaha dan branding", 1200000, "Sekali", None, None),
              ("2026-11", ASET, "Mesin cuci dan pengering", 24000000, "Sekali", None, 60),
              ("2026-11", MODAL, "Modal awal", 30000000, "Sekali", None, None)],
        period=("2026-11", "2027-10"),
        tx=[("2026-10-31", "Cuci setrika Oktober", JUAL, 1400 * 8000, None, "Cuci setrika", "1400", None, None),
            ("2026-10-31", "Cuci kilat Oktober", JUAL, 250 * 12000, None, "Cuci kilat", "250", None, None),
            ("2026-10-20", "Cuci sprei hotel (borongan)", JUAL, 300000, None, None, None, None, None),
            ("2026-10-05", "Deterjen, plastik, listrik", BAHAN, None, 4600000, None, None, None, None),
            ("2026-10-30", "Gaji karyawan", GAJI, None, 3600000, None, None, None, None),
            ("2026-10-01", "Sewa ruko", SEWA, None, 2000000, None, None, None, None),
            ("2026-10-01", "Mesin cuci dan pengering", ASET, None, 24000000, None, None, None, 60),
            ("2026-10-01", "Modal awal", MODAL, 30000000, None, None, None, None, None)],
    ),
    "C_cabai": dict(
        jenis="Pertanian (budidaya cabai satu musim, 4 bulan)",
        products=[("Cabai merah", "kg", 35000, 250, [("Panen dan angkut", 3000, 1)])],
        plan=[(None, GAJI, "Upah harian", 2000000, "Bulanan", None, None),
              ("2026-11", SEWA, "Sewa lahan semusim", 4000000, "Sekali", None, None),
              ("2026-11", BAHAN, "Bibit", 2500000, "Sekali", None, None),
              ("2026-12", BAHAN, "Pupuk dan obat", 6000000, "Sekali", None, None),
              ("2026-11", LAIN, "Mulsa dan ajir", 3000000, "Sekali", None, None),
              ("2027-02", TRANS, "Transport panen", 800000, "Sekali", None, None),
              ("2026-11", ASET, "Pompa dan sprayer", 3600000, "Sekali", None, 24)],
        tx=[("2026-10-10", "Sewa lahan semusim", SEWA, None, 4000000, None, None, None, None),
            ("2026-10-12", "Bibit", BAHAN, None, 2500000, None, None, None, None),
            ("2026-10-12", "Mulsa dan ajir", LAIN, None, 3000000, None, None, None, None),
            ("2026-10-15", "Pompa dan sprayer", ASET, None, 3600000, None, None, None, 24)],
    ),
    "D_rental_traktor": dict(
        jenis="Sewa alat (rental traktor, aset besar)",
        products=[("Sewa traktor", "hari", 600000, 18, [("Solar", 150000, 1), ("Upah operator", 100000, 1)])],
        plan=[(None, LAIN, "Perawatan dan servis", 750000, "Bulanan", None, None),
              ("2026-11", ASET, "Traktor roda 4", 90000000, "Sekali", None, 60),
              ("2026-11", MODAL, "Modal pemilik", 50000000, "Sekali", None, None)],
        tx=[("2026-10-31", "Sewa traktor Oktober", JUAL, 16 * 600000, None, "Sewa traktor", "16", None, None),
            ("2026-10-31", "Solar", TRANS, None, 2500000, None, None, None, None),
            ("2026-10-31", "Upah operator", GAJI, None, 1600000, None, None, "Variabel", None),
            ("2026-10-15", "Servis rutin", LAIN, None, 750000, None, None, None, None),
            ("2026-10-01", "Traktor roda 4", ASET, None, 90000000, None, None, None, 60),
            ("2026-10-01", "Modal pemilik", MODAL, 50000000, None, None, None, None, None)],
    ),
}

def months(sim):
    if "period" in sim:
        a, b = sim["period"]
    else:
        dated = sorted(m for m, *_ in sim["plan"] if m)
        a, b = dated[0], dated[-1]
    ya, ma = map(int, a.split("-")); yb, mb = map(int, b.split("-"))
    return (yb - ya) * 12 + (mb - ma) + 1

def expected_plan(sim):
    n = months(sim)
    fixed = 0; other_var = 0; other_rev = 0; dep = 0
    for m, cat, name, amt, rec, nat, life in sim["plan"]:
        per_month = amt if rec == "Bulanan" else D(amt) / n
        if cat == ASET: dep += r0(D(amt) / life)
        elif cat in NATURE:
            if (nat or NATURE[cat]) in ("V", "Variabel"): other_var += r0(per_month)
            else: fixed += r0(per_month)
        elif cat in (JUAL, LAIN_IN): other_rev += r0(per_month)
    fc = fixed + dep
    base = fc + other_var - other_rev
    prods = []
    for name, unit, price, target, comps in sim["products"]:
        vc = sum(D(c) / D(y) for _, c, y in comps)
        prods.append((name, unit, price, D(target), vc, D(price) - vc))
    total = sum(p[3] for p in prods)
    wcm = sum(p[5] * p[3] for p in prods) / total
    bep_units = D(base) / wcm
    rev = sum(r0(p[3] * p[2]) for p in prods) + other_rev
    var = sum(r0(p[3] * p[4]) for p in prods) + other_var
    tc = fc + var
    return dict(months=n, fixed_total=fc, bep_base=base,
                bep_units={p[0]: str((bep_units * p[3] / total).quantize(D("0.001"), ROUND_HALF_UP)) for p in prods},
                bep_rupiah=sum(r0(D(base) * (p[3] / total) * p[2] / wcm) for p in prods),
                vc={p[0]: r0(p[4]) for p in prods},
                revenue=rev, total_cost=tc, profit=rev - tc,
                rc=str((D(rev) / tc).quantize(D("0.01"), ROUND_HALF_UP)),
                bc=str((D(rev - tc) / tc).quantize(D("0.01"), ROUND_HALF_UP)))

def expected_actual(sim):
    price = {p[0]: p[2] for p in sim["products"]}
    vc = {p[0]: sum(D(c) / D(y) for _, c, y in p[4]) for p in sim["products"]}
    target = {p[0]: D(p[3]) for p in sim["products"]}
    units = {}; revenue = 0; unnamed = 0; fixed = 0; variable = 0; dep = 0
    for date, desc, cat, inc, exp, prod, qty, nat, life in sim["tx"]:
        if inc and cat in (JUAL, LAIN_IN):
            revenue += inc
            if prod: units[prod] = units.get(prod, 0) + D(qty)
            else: unnamed += inc
        elif exp and cat in NATURE:
            if (nat or NATURE[cat]) in ("V", "Variabel"): variable += exp
            else: fixed += exp
        elif exp and cat == ASET:
            dep += exp // life
    planned_rev = sum(target[p] * price[p] for p in price)
    for p in price:
        if unnamed:
            units[p] = units.get(p, 0) + (D(unnamed) * target[p] / planned_rev).quantize(D("0.001"), ROUND_HALF_UP)
    fc = fixed + dep
    std = sum(r0(q * vc[p]) for p, q in units.items())
    urev = sum(q * price[p] for p, q in units.items()); ucm = sum(q * (price[p] - vc[p]) for p, q in units.items())
    bep = r0(D(fc) * urev / ucm) if units and ucm > 0 else None
    tc = fc + variable
    return dict(revenue=revenue, units={p: str(q) for p, q in units.items()}, fixed_and_dep=fc, bep=bep,
                reached=None if bep is None else revenue >= bep, variance=variable - std,
                total_cost=tc, profit=revenue - tc, rc=str((D(revenue) / tc).quantize(D("0.01"), ROUND_HALF_UP)) if tc else None)

def write(name, sim):
    wb = Workbook(); ws = wb.active; ws.title = "Petunjuk"
    ws.append(["Kunci", "Nilai"]); ws.append(["Format", "MANKEU_TEMPLATE"])
    t = wb.create_sheet("Transaksi"); t.append(TX)
    for r in sim["tx"]: t.append(list(r))
    b = wb.create_sheet("Rencana Anggaran"); b.append(BUDGET)
    for m, cat, item, amt, rec, nat, life in sim["plan"]:
        b.append([m, cat, item, None, None, None, amt, rec, nat, life])
    p = wb.create_sheet("Produk"); p.append(PRODUCT)
    rc = wb.create_sheet("Resep Produk"); rc.append(RECIPE)
    for pname, unit, price, target, comps in sim["products"]:
        p.append([pname, unit, price, target])
        for c, cost, y in comps: rc.append([pname, c, cost, y])
    s = wb.create_sheet("Stok Barang"); s.append(STOCK)
    for r in sim.get("stock", []): s.append(list(r))
    wb.save(os.path.join(OUT, f"{name}.xlsx"))

expected = {}
for name, sim in SIMS.items():
    write(name, sim)
    expected[name] = dict(jenis=sim["jenis"], plan=expected_plan(sim), actual=expected_actual(sim))
json.dump(expected, open(os.path.join(OUT, "expected.json"), "w"), indent=1, ensure_ascii=False)
print(json.dumps(expected, indent=1, ensure_ascii=False))
