"""Schema 10 simulations: seasonal sales, per-unit stock, period undo on batch reversal.

Writes Mankeu template files (F as v4 with the selling-month column, G/H as v3 without it) and
expected figures computed here from first principles, not by calling the app:
- break-even = the uniform volume scale at which the period's profit is exactly zero;
- capital need = deepest cumulative cash deficit of operations + investment, walked month by month.

Run: python -I make_sims.py  (needs openpyxl). The .xlsx files are generated, not committed.

Historical record (2026-10-29, phase-1 cleanup P8): this run predates template v5 (DEC-028); the
current importer rejects these files by their headers. Kept unchanged as the evidence of that run;
the v5 generator is `docs/evidence/2026-10-22_simulasi_skala_umkm/make_sims.py`.
"""
from decimal import Decimal as D, ROUND_HALF_UP
import json, os
from openpyxl import Workbook

OUT = os.path.dirname(os.path.abspath(__file__))
TX = ["Tanggal", "Keterangan", "Kategori", "Pemasukan (Rp)", "Pengeluaran (Rp)", "Produk", "Jumlah", "Sifat Biaya", "Umur Pakai (bulan)"]
BUDGET = ["Bulan", "Kategori", "Item", "Jumlah", "Satuan", "Harga Satuan (Rp)", "Total (Rp)", "Berulang", "Sifat Biaya", "Umur Pakai (bulan)"]
PRODUCT_V3 = ["Nama", "Satuan", "Harga Jual (Rp)", "Target per Bulan"]
PRODUCT_V4 = PRODUCT_V3 + ["Mulai Dijual (bulan ke-)"]
RECIPE = ["Produk", "Komponen", "Biaya (Rp)", "Untuk Berapa Unit"]
STOCK = ["Nama", "Satuan", "Harga Beli (Rp)", "Jumlah Stok", "Batas Minimum"]

JUAL = "Penjualan & Jasa"
BAHAN, GAJI, SEWA, LISTRIK = "Bahan Baku & Stok", "Gaji & Upah", "Sewa Tempat", "Listrik, Air & Internet"
MODAL, ASET = "Setoran Modal Pemilik", "Beli Peralatan & Aset"
VARIABLE = {BAHAN}  # default cost nature of the basic categories used here

def r0(x): return int(D(x).quantize(D(1), ROUND_HALF_UP))

SIMS = {
    # Seasonal: chilli is harvested from month 4, seedlings sell from month 1. Mixed seasons in one mix.
    "F_cabai_musiman": dict(
        jenis="Pertanian 12 bulan: cabai panen mulai bulan ke-4 + jual bibit sejak bulan 1",
        period=("2026-11", "2027-10"),
        # name, unit, price, target/month, sell from month, recipe [(component, cost, for how many units)]
        products=[("Cabai merah", "kg", 35000, 400, 4, [("Panen dan angkut", 3000, 1), ("Karung", 5000, 50)]),
                  ("Bibit cabai polybag", "batang", 1500, 2000, 1, [("Polybag dan media", 30000, 100), ("Benih", 200, 1)])],
        # month, category, item, total, recurrence, nature, useful life
        plan=[(None, GAJI, "Upah harian", 2500000, "Bulanan", None, None),
              ("2026-11", SEWA, "Sewa lahan setahun", 12000000, "Sekali", None, None),
              (None, BAHAN, "Pupuk dan obat", 1500000, "Bulanan", None, None),
              ("2026-11", ASET, "Pompa dan sprayer", 3600000, "Sekali", None, 24),
              ("2026-11", MODAL, "Modal awal", 20000000, "Sekali", None, None)],
        tx=[("2026-11-01", "Modal awal", MODAL, 20000000, None, None, None, None, None),
            ("2026-11-02", "Sewa lahan setahun", SEWA, None, 12000000, None, None, None, None),
            ("2026-11-05", "Pompa dan sprayer", ASET, None, 3600000, None, None, None, 24),
            ("2026-11-20", "Pupuk dan obat", BAHAN, None, 1500000, None, None, None, None),
            ("2026-11-30", "Upah harian", GAJI, None, 2500000, None, None, None, None),
            ("2026-11-30", "Jual bibit November", JUAL, 1800 * 1500, None, "Bibit cabai polybag", "1800", None, None)],
        stock=[("Benih cabai", "pack", 85000, "4", "1")],
        v4=True,
    ),
    # Two branches of one shop. Each holds its own "Gula"; "Kantong plastik" is added by hand as shared.
    "G_toko_pusat": dict(
        jenis="Dagang, cabang pusat (berkas v3 tanpa kolom bulan mulai)",
        period=("2026-11", "2027-04"),
        products=[("Beras 5 kg", "karung", 75000, 120, 1, [("Harga beli", 65000, 1)])],
        plan=[(None, SEWA, "Sewa kios", 1500000, "Bulanan", None, None),
              (None, LISTRIK, "Listrik", 300000, "Bulanan", None, None)],
        tx=[("2026-11-01", "Modal awal", MODAL, 10000000, None, None, None, None, None),
            ("2026-11-30", "Jual beras November", JUAL, 100 * 75000, None, "Beras 5 kg", "100", None, None),
            ("2026-11-03", "Kulakan beras", BAHAN, None, 6500000, None, None, None, None)],
        stock=[("Gula", "kg", 15500, "50", "10"), ("Kantong plastik", "pak", 12000, "9", "2")],
    ),
    "H_toko_cabang": dict(
        jenis="Dagang, cabang 2 (berkas v3)",
        period=("2026-11", "2027-04"),
        products=[("Beras 5 kg", "karung", 76000, 80, 1, [("Harga beli", 65000, 1)])],
        plan=[(None, SEWA, "Sewa kios", 1000000, "Bulanan", None, None)],
        tx=[("2026-11-30", "Jual beras November", JUAL, 70 * 76000, None, "Beras 5 kg", "70", None, None)],
        stock=[("Gula", "kg", 15800, "20", "5"), ("Kantong plastik", "pak", 12000, "3", "1")],
    ),
}

def month_list(sim):
    (ya, ma), (yb, mb) = (map(int, x.split("-")) for x in sim["period"])
    n = (yb - ya) * 12 + (mb - ma) + 1
    return [f"{ya + (ma - 1 + i) // 12}-{(ma - 1 + i) % 12 + 1:02d}" for i in range(n)]

def vc(comps): return sum(D(c) / D(y) for _, c, y in comps)

def expected_plan(sim):
    ms = month_list(sim); n = len(ms)
    def selling(p, i): return i + 1 >= p[4]
    # Month by month cash: what the owner actually lives through.
    operating, investing, financing = [], [], []
    for i, m in enumerate(ms):
        inc = sum(D(p[3]) * p[2] for p in sim["products"] if selling(p, i))
        exp = sum(D(p[3]) * vc(p[5]) for p in sim["products"] if selling(p, i))
        inv = fin = 0
        for pm, cat, _, amt, rec, _, _ in sim["plan"]:
            hits = rec == "Bulanan" or (pm or ms[0]) == m
            if not hits: continue
            if cat == ASET: inv -= amt
            elif cat == MODAL: fin += amt
            elif cat == JUAL: inc += amt
            else: exp += amt
        operating.append(r0(inc) - r0(exp)); investing.append(inv); financing.append(fin)
    cum = 0; need = 0
    for o, v in zip(operating, investing):
        cum += o + v; need = max(need, -cum)
    cum = 0; worst = 0
    for o, v, f in zip(operating, investing, financing):
        cum += o + v + f; worst = min(worst, cum)
    # Profit per plan month (depreciation in, asset purchase out), and BEP by scaling all volumes.
    dep = sum(r0(D(amt) / life) for _, cat, _, amt, _, _, life in sim["plan"] if cat == ASET)
    fixed = sum((amt if rec == "Bulanan" else r0(D(amt) / n)) for _, cat, _, amt, rec, nat, _ in sim["plan"]
                if cat not in (ASET, MODAL, JUAL) and cat not in VARIABLE)
    other_var = sum((amt if rec == "Bulanan" else r0(D(amt) / n)) for _, cat, _, amt, rec, nat, _ in sim["plan"] if cat in VARIABLE)
    base = fixed + dep + other_var
    s = {p[0]: max(0, n - p[4] + 1) for p in sim["products"]}
    margin_period = sum(D(p[3]) * s[p[0]] * (p[2] - vc(p[5])) for p in sim["products"])  # at target volume
    k = D(base) * n / margin_period  # volume scale with zero period profit
    revenue = r0(sum(D(p[3]) * s[p[0]] * p[2] for p in sim["products"]) / n)
    variable = r0(sum(D(p[3]) * s[p[0]] * vc(p[5]) for p in sim["products"]) / n) + other_var
    total_cost = fixed + dep + variable
    return dict(
        months=n, schedule_operating=operating, capital_need=need, capital_planned=sum(financing), capital_shortfall=-worst,
        fixed_total=fixed + dep, bep_base=base,
        bep_units_per_selling_month={p[0]: str((k * p[3]).quantize(D("0.001"), ROUND_HALF_UP)) for p in sim["products"]},
        bep_rupiah_per_month=r0(k * sum(D(p[3]) * s[p[0]] * p[2] for p in sim["products"]) / n),
        revenue_per_month=revenue, total_cost_per_month=total_cost, profit_per_month=revenue - total_cost,
        period_profit_cash=sum(operating),
        rc=str((D(revenue) / total_cost).quantize(D("0.01"), ROUND_HALF_UP)),
    )

def write(name, sim):
    wb = Workbook(); ws = wb.active; ws.title = "Petunjuk"
    ws.append(["Kunci", "Nilai"]); ws.append(["Format", "MANKEU_TEMPLATE"])
    ws.append(["Periode rencana", f"{sim['period'][0]} s.d. {sim['period'][1]}"])
    t = wb.create_sheet("Transaksi"); t.append(TX)
    for r in sim["tx"]: t.append(list(r))
    b = wb.create_sheet("Rencana Anggaran"); b.append(BUDGET)
    for m, cat, item, amt, rec, nat, life in sim["plan"]:
        b.append([m, cat, item, None, None, None, amt, rec, nat, life])
    p = wb.create_sheet("Produk"); p.append(PRODUCT_V4 if sim.get("v4") else PRODUCT_V3)
    rc = wb.create_sheet("Resep Produk"); rc.append(RECIPE)
    for pname, unit, price, target, start, comps in sim["products"]:
        p.append([pname, unit, price, target] + ([start] if sim.get("v4") else []))
        for c, cost, y in comps: rc.append([pname, c, cost, y])
    s = wb.create_sheet("Stok Barang"); s.append(STOCK)
    for r in sim.get("stock", []): s.append(list(r))
    wb.save(os.path.join(OUT, f"{name}.xlsx"))

def batch_money(sim):
    return dict(incoming=sum(r[3] or 0 for r in sim["tx"]), outgoing=sum(r[4] or 0 for r in sim["tx"]))

expected = {}
for name, sim in SIMS.items():
    write(name, sim)
    expected[name] = dict(jenis=sim["jenis"], plan=expected_plan(sim), batch=batch_money(sim))
# The same chilli plan if it sold from month 1: what the app used to show (capital need understated).
flat = dict(SIMS["F_cabai_musiman"], products=[p[:4] + (1,) + p[5:] for p in SIMS["F_cabai_musiman"]["products"]])
expected["F_cabai_jika_jual_sejak_bulan_1"] = dict(plan=expected_plan(flat))
json.dump(expected, open(os.path.join(OUT, "expected.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
print(json.dumps(expected, indent=1, ensure_ascii=False))
